import os
import io
import csv
import shutil
import tempfile
from datetime import datetime
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Depends, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import func, and_

from database import get_db, Document, Item, DOC_TYPES
from extractor import (
    extract_text, detect_doc_type, extract_with_claude,
    search_market_price, generate_estimate
)

app = FastAPI(title="ビジネス書類データベース", version="2.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"],
                   allow_methods=["*"], allow_headers=["*"])

ALLOWED_EXT = {"pdf", "xlsx", "xls", "docx", "jpg", "jpeg", "png", "webp"}


# ── Pydantic モデル ──────────────────────────────

class EstimateRequest(BaseModel):
    client_name: str
    company_name: str
    memo: str = ""
    items: list

class MarketRequest(BaseModel):
    item_name: str
    item_id: Optional[int] = None

class StatusUpdate(BaseModel):
    status: str

class MemoUpdate(BaseModel):
    memo: str


# ── アップロード ─────────────────────────────────

@app.post("/api/upload")
async def upload_document(
    file: UploadFile = File(...),
    doc_type: str = "auto",
    db: Session = Depends(get_db)
):
    ext = file.filename.lower().rsplit(".", 1)[-1]
    if ext not in ALLOWED_EXT:
        raise HTTPException(400, f"非対応の形式: {file.filename}")

    with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
        shutil.copyfileobj(file.file, tmp)
        tmp_path = tmp.name

    try:
        text = extract_text(tmp_path, file.filename)
        if not text.strip():
            raise HTTPException(400, "テキストを抽出できませんでした")

        # 書類種別の判定
        resolved_type = detect_doc_type(text, file.filename) if doc_type == "auto" else doc_type

        # AI抽出
        try:
            data = extract_with_claude(text, file.filename, resolved_type)
        except Exception as e:
            raise HTTPException(500, f"AI抽出エラー: {e}")

        # DB保存
        doc = Document(
            doc_type    = resolved_type,
            doc_number  = data.get("doc_number", ""),
            filename    = file.filename,
            client_name = data.get("client_name", ""),
            doc_date    = data.get("doc_date", ""),
            due_date    = data.get("due_date", ""),
            total_amount= float(data.get("total_amount", 0) or 0),
            tax_amount  = float(data.get("tax_amount", 0) or 0),
            currency    = data.get("currency", "JPY"),
            memo        = data.get("memo", ""),
            raw_text    = text[:10000],
        )
        db.add(doc)
        db.flush()

        for it in data.get("items", []):
            db.add(Item(
                document_id = doc.id,
                name        = str(it.get("name", "")),
                quantity    = float(it.get("quantity", 1) or 1),
                unit        = str(it.get("unit", "")),
                unit_price  = float(it.get("unit_price", 0) or 0),
                total_price = float(it.get("total_price", 0) or 0),
                note        = str(it.get("note", "")),
            ))

        db.commit()
        db.refresh(doc)
        return {
            "id": doc.id,
            "doc_type": doc.doc_type,
            "filename": doc.filename,
            "client_name": doc.client_name,
            "total_amount": doc.total_amount,
            "items_count": len(data.get("items", [])),
        }
    finally:
        os.unlink(tmp_path)


# ── 書類一覧・詳細 ────────────────────────────────

@app.get("/api/documents")
def list_documents(
    doc_type: str = "",
    client: str = "",
    status: str = "",
    db: Session = Depends(get_db)
):
    q = db.query(Document)
    if doc_type:
        q = q.filter(Document.doc_type == doc_type)
    if client:
        q = q.filter(Document.client_name.contains(client))
    if status:
        q = q.filter(Document.status == status)
    docs = q.order_by(Document.uploaded_at.desc()).all()
    return [_doc_summary(d) for d in docs]


@app.get("/api/documents/{doc_id}")
def get_document(doc_id: int, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404, "書類が見つかりません")
    return _doc_detail(doc)


@app.patch("/api/documents/{doc_id}/status")
def update_status(doc_id: int, body: StatusUpdate, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404)
    doc.status = body.status
    db.commit()
    return {"ok": True}


@app.patch("/api/documents/{doc_id}/memo")
def update_memo(doc_id: int, body: MemoUpdate, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404)
    doc.memo = body.memo
    db.commit()
    return {"ok": True}


@app.delete("/api/documents/{doc_id}")
def delete_document(doc_id: int, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404)
    db.delete(doc)
    db.commit()
    return {"ok": True}


# ── 品目検索 ─────────────────────────────────────

@app.get("/api/search")
def search_items(q: str = Query(..., min_length=1), db: Session = Depends(get_db)):
    items = (db.query(Item).join(Document)
             .filter(Item.name.contains(q))
             .order_by(Item.total_price.desc()).all())
    return [_item_row(i) for i in items]


# ── 相場価格検索 ──────────────────────────────────

@app.post("/api/market-price")
async def get_market_price(body: MarketRequest, db: Session = Depends(get_db)):
    try:
        result = search_market_price(body.item_name)
    except Exception as e:
        raise HTTPException(500, f"相場検索エラー: {e}")

    # DBに保存（item_idがある場合）
    if body.item_id:
        item = db.query(Item).filter(Item.id == body.item_id).first()
        if item:
            item.market_low  = float(result.get("market_low", 0) or 0)
            item.market_high = float(result.get("market_high", 0) or 0)
            item.market_memo = str(result.get("memo", ""))
            db.commit()

    return result


# ── 統計・ダッシュボード ──────────────────────────

@app.get("/api/stats")
def get_stats(db: Session = Depends(get_db)):
    total_docs  = db.query(func.count(Document.id)).scalar()
    total_items = db.query(func.count(Item.id)).scalar()
    total_amt   = db.query(func.sum(Document.total_amount)).scalar() or 0

    # 書類種別ごとの件数・金額
    by_type = (
        db.query(Document.doc_type,
                 func.count(Document.id).label("count"),
                 func.sum(Document.total_amount).label("amount"))
        .group_by(Document.doc_type).all()
    )

    # 月別金額推移（直近12ヶ月）
    monthly = (
        db.query(
            func.strftime("%Y-%m", Document.doc_date).label("month"),
            func.sum(Document.total_amount).label("amount"),
            func.count(Document.id).label("count"),
        )
        .filter(Document.doc_date != "")
        .group_by(func.strftime("%Y-%m", Document.doc_date))
        .order_by(func.strftime("%Y-%m", Document.doc_date).desc())
        .limit(12).all()
    )

    # 取引先ランキング
    clients = (
        db.query(Document.client_name,
                 func.count(Document.id).label("count"),
                 func.sum(Document.total_amount).label("amount"))
        .filter(Document.client_name != "")
        .group_by(Document.client_name)
        .order_by(func.sum(Document.total_amount).desc())
        .limit(10).all()
    )

    # よく出る品目
    top_items = (
        db.query(Item.name,
                 func.count(Item.id).label("count"),
                 func.avg(Item.unit_price).label("avg_price"))
        .group_by(Item.name)
        .order_by(func.count(Item.id).desc())
        .limit(10).all()
    )

    return {
        "total_docs":  total_docs,
        "total_items": total_items,
        "total_amount": total_amt,
        "by_type": [{"type": r.doc_type, "count": r.count, "amount": r.amount or 0} for r in by_type],
        "monthly":  [{"month": r.month, "amount": r.amount or 0, "count": r.count} for r in reversed(monthly)],
        "clients":  [{"name": r.client_name, "count": r.count, "amount": r.amount or 0} for r in clients],
        "top_items":[{"name": r.name, "count": r.count, "avg_price": round(r.avg_price or 0)} for r in top_items],
    }


# ── 取引先分析 ────────────────────────────────────

@app.get("/api/clients")
def get_clients(db: Session = Depends(get_db)):
    rows = (
        db.query(Document.client_name,
                 func.count(Document.id).label("count"),
                 func.sum(Document.total_amount).label("total"),
                 func.max(Document.doc_date).label("last_date"))
        .filter(Document.client_name != "")
        .group_by(Document.client_name)
        .order_by(func.sum(Document.total_amount).desc()).all()
    )
    return [{"name": r.client_name, "count": r.count,
             "total": r.total or 0, "last_date": r.last_date or ""} for r in rows]


@app.get("/api/clients/{client_name}/documents")
def get_client_documents(client_name: str, db: Session = Depends(get_db)):
    docs = (db.query(Document)
            .filter(Document.client_name == client_name)
            .order_by(Document.doc_date.desc()).all())
    return [_doc_summary(d) for d in docs]


# ── CSVエクスポート ───────────────────────────────

@app.get("/api/export/csv")
def export_csv(
    doc_type: str = "",
    client: str = "",
    db: Session = Depends(get_db)
):
    q = db.query(Document)
    if doc_type:
        q = q.filter(Document.doc_type == doc_type)
    if client:
        q = q.filter(Document.client_name.contains(client))
    docs = q.order_by(Document.doc_date.desc()).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["書類種別", "書類番号", "ファイル名", "取引先", "日付",
                     "支払期限/納期", "合計金額", "消費税", "ステータス", "アップロード日"])
    for d in docs:
        writer.writerow([d.doc_type, d.doc_number, d.filename, d.client_name,
                         d.doc_date, d.due_date, d.total_amount, d.tax_amount,
                         d.status, d.uploaded_at.strftime("%Y-%m-%d") if d.uploaded_at else ""])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=documents.csv"}
    )


@app.get("/api/export/items-csv")
def export_items_csv(db: Session = Depends(get_db)):
    items = db.query(Item).join(Document).order_by(Document.doc_date.desc()).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["書類種別", "取引先", "日付", "品目名", "数量", "単位",
                     "単価", "合計", "相場下限", "相場上限", "備考"])
    for i in items:
        writer.writerow([
            i.document.doc_type, i.document.client_name, i.document.doc_date,
            i.name, i.quantity, i.unit, i.unit_price, i.total_price,
            i.market_low or "", i.market_high or "", i.note
        ])
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=items.csv"}
    )


# ── 見積書自動作成 ────────────────────────────────

@app.post("/api/generate-estimate")
async def create_estimate(body: EstimateRequest):
    try:
        data = generate_estimate(body.client_name, body.items,
                                 body.company_name, body.memo)
        return data
    except Exception as e:
        raise HTTPException(500, f"見積書生成エラー: {e}")


# ── ヘルパー ─────────────────────────────────────

def _doc_summary(d: Document) -> dict:
    return {
        "id": d.id, "doc_type": d.doc_type, "doc_number": d.doc_number,
        "filename": d.filename, "client_name": d.client_name,
        "doc_date": d.doc_date, "due_date": d.due_date,
        "total_amount": d.total_amount, "tax_amount": d.tax_amount,
        "status": d.status, "items_count": len(d.items),
        "uploaded_at": d.uploaded_at.isoformat() if d.uploaded_at else "",
    }


def _doc_detail(d: Document) -> dict:
    s = _doc_summary(d)
    s["memo"] = d.memo
    s["items"] = [{
        "id": i.id, "name": i.name, "quantity": i.quantity,
        "unit": i.unit, "unit_price": i.unit_price, "total_price": i.total_price,
        "note": i.note, "market_low": i.market_low, "market_high": i.market_high,
        "market_memo": i.market_memo,
    } for i in d.items]
    return s


def _item_row(i: Item) -> dict:
    return {
        "item_id": i.id, "name": i.name, "quantity": i.quantity,
        "unit": i.unit, "unit_price": i.unit_price, "total_price": i.total_price,
        "note": i.note, "market_low": i.market_low, "market_high": i.market_high,
        "doc_id": i.document_id, "doc_type": i.document.doc_type,
        "client_name": i.document.client_name, "doc_date": i.document.doc_date,
        "filename": i.document.filename,
    }


app.mount("/static", StaticFiles(directory="static"), name="static")
app.mount("/", StaticFiles(directory="static", html=True), name="root")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)


# ── Excel 書類出力（既存書類） ─────────────────────────

@app.get("/api/export/excel/{doc_id}")
def export_excel_doc(doc_id: int, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404, "書類が見つかりません")
    from build_templates import create_estimate_sheet, create_invoice_sheet
    from openpyxl import Workbook
    import io
    wb = Workbook()
    del wb["Sheet"]
    info = {
        "client": doc.client_name, "doc_date": doc.doc_date,
        "subject": doc.memo or "", "delivery_date": doc.due_date or "ー",
        "delivery_place": "ー", "payment_method": "従来通り", "staff": "木戸　志朗",
    }
    items = [{"name": i.name, "qty": i.quantity, "unit": i.unit,
              "price": i.unit_price, "note": i.note} for i in doc.items]
    if doc.doc_type == "請求書":
        create_invoice_sheet(wb, items=items, info=info)
        fname = f"請求書_{doc.client_name}_{doc.doc_date}.xlsx"
    else:
        create_estimate_sheet(wb, items=items, info=info)
        fname = f"見積書_{doc.client_name}_{doc.doc_date}.xlsx"
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    from urllib.parse import quote
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ── Excel 書類出力（見積書作成画面から） ──────────────

class ExcelGenRequest(BaseModel):
    client_name: str
    company_name: str = "アースワーク株式会社"
    memo: str = ""
    doc_type: str = "見積書"
    items: list

@app.post("/api/export/excel/generate")
async def generate_excel(body: ExcelGenRequest):
    from build_templates import create_estimate_sheet, create_invoice_sheet
    from openpyxl import Workbook
    import io, datetime
    wb = Workbook()
    del wb["Sheet"]
    info = {
        "client": body.client_name,
        "doc_date": datetime.date.today().strftime("%Y年%m月%d日"),
        "subject": body.memo, "delivery_date": "ー", "staff": "木戸　志朗",
    }
    if body.doc_type == "請求書":
        create_invoice_sheet(wb, items=body.items, info=info)
        fname = f"請求書_{body.client_name}.xlsx"
    else:
        create_estimate_sheet(wb, items=body.items, info=info)
        fname = f"見積書_{body.client_name}.xlsx"
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    from urllib.parse import quote
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})
